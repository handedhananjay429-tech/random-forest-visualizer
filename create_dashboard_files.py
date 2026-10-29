import pandas as pd
import numpy as np
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, PieChart, Reference, Series
from openpyxl.chart.label import DataLabelList

def generate_dashboard_files():
    print("Generating transactional data for 18,450 orders...")
    np.random.seed(42)

    # Months, Sales (Lakh), Profit (Lakh), Orders
    months_data = [
        {"Month": "Jan", "MonthNum": 1, "Quarter": "Q1", "Days": 31, "Sales_Lakh": 42.0, "Profit_Lakh": 8.0, "Orders": 1120},
        {"Month": "Feb", "MonthNum": 2, "Quarter": "Q1", "Days": 29, "Sales_Lakh": 45.0, "Profit_Lakh": 9.0, "Orders": 1180},
        {"Month": "Mar", "MonthNum": 3, "Quarter": "Q1", "Days": 31, "Sales_Lakh": 51.0, "Profit_Lakh": 11.0, "Orders": 1320},
        {"Month": "Apr", "MonthNum": 4, "Quarter": "Q2", "Days": 30, "Sales_Lakh": 48.0, "Profit_Lakh": 9.0, "Orders": 1260},
        {"Month": "May", "MonthNum": 5, "Quarter": "Q2", "Days": 31, "Sales_Lakh": 55.0, "Profit_Lakh": 12.0, "Orders": 1410},
        {"Month": "Jun", "MonthNum": 6, "Quarter": "Q2", "Days": 30, "Sales_Lakh": 60.0, "Profit_Lakh": 14.0, "Orders": 1530},
        {"Month": "Jul", "MonthNum": 7, "Quarter": "Q3", "Days": 31, "Sales_Lakh": 58.0, "Profit_Lakh": 13.0, "Orders": 1480},
        {"Month": "Aug", "MonthNum": 8, "Quarter": "Q3", "Days": 31, "Sales_Lakh": 63.0, "Profit_Lakh": 15.0, "Orders": 1600},
        {"Month": "Sep", "MonthNum": 9, "Quarter": "Q3", "Days": 30, "Sales_Lakh": 67.0, "Profit_Lakh": 16.0, "Orders": 1710},
        {"Month": "Oct", "MonthNum": 10, "Quarter": "Q4", "Days": 31, "Sales_Lakh": 70.0, "Profit_Lakh": 17.0, "Orders": 1780},
        {"Month": "Nov", "MonthNum": 11, "Quarter": "Q4", "Days": 30, "Sales_Lakh": 76.0, "Profit_Lakh": 19.0, "Orders": 1940},
        {"Month": "Dec", "MonthNum": 12, "Quarter": "Q4", "Days": 31, "Sales_Lakh": 85.0, "Profit_Lakh": 22.0, "Orders": 2120},
    ]

    # Category proportions of total sales (Total 7.2 Cr = 720 Lakh):
    # Electronics: 2.4 Cr (33.33%), Furniture: 1.7 Cr (23.61%), Clothing: 1.3 Cr (18.06%), Grocery: 1.1 Cr (15.28%), Toys: 0.7 Cr (9.72%)
    categories = ["Electronics", "Furniture", "Clothing", "Grocery", "Toys"]
    cat_weights = [2.4/7.2, 1.7/7.2, 1.3/7.2, 1.1/7.2, 0.7/7.2]

    # Region weights: West: 38%, South: 27%, North: 20%, East: 15%
    regions = ["West", "South", "North", "East"]
    reg_weights = [0.38, 0.27, 0.20, 0.15]

    segments = ["Consumer", "Corporate", "Home Office"]
    seg_weights = [0.55, 0.30, 0.15]

    payment_modes = ["UPI", "Credit Card", "Net Banking", "Debit Card", "Cash on Delivery"]
    pay_weights = [0.45, 0.25, 0.15, 0.10, 0.05]

    subcats = {
        "Electronics": ["Smartphones", "Laptops", "Accessories", "Audio", "Tablets"],
        "Furniture": ["Chairs", "Desks", "Bookcases", "Sofas", "Storage"],
        "Clothing": ["Men's Apparel", "Women's Apparel", "Footwear", "Sportswear"],
        "Grocery": ["Packaged Food", "Beverages", "Snacks", "Gourmet", "Personal Care"],
        "Toys": ["Board Games", "Action Figures", "Puzzles", "Educational Toys", "Outdoor"]
    }

    all_rows = []
    order_counter = 1

    for m in months_data:
        m_orders = m["Orders"]
        target_sales = m["Sales_Lakh"] * 100000.0
        target_profit = m["Profit_Lakh"] * 100000.0

        # Generate base random sales weights
        raw_weights = np.random.pareto(a=2.5, size=m_orders) + 0.1
        sales_vals = (raw_weights / raw_weights.sum()) * target_sales

        # Generate profit with minor noise around target margin
        target_margin = target_profit / target_sales
        profit_factors = np.random.normal(loc=target_margin, scale=0.03, size=m_orders)
        profit_factors = np.clip(profit_factors, 0.05, 0.45)
        raw_profits = sales_vals * profit_factors
        profit_vals = (raw_profits / raw_profits.sum()) * target_profit

        # Assign categories and regions based on distribution
        m_categories = np.random.choice(categories, size=m_orders, p=cat_weights)
        m_regions = np.random.choice(regions, size=m_orders, p=reg_weights)
        m_segments = np.random.choice(segments, size=m_orders, p=seg_weights)
        m_payments = np.random.choice(payment_modes, size=m_orders, p=pay_weights)

        days = np.random.randint(1, m["Days"] + 1, size=m_orders)

        for i in range(m_orders):
            cat = m_categories[i]
            subcat = np.random.choice(subcats[cat])
            dt = datetime.date(2024, m["MonthNum"], int(days[i]))
            s_val = round(float(sales_vals[i]), 2)
            p_val = round(float(profit_vals[i]), 2)
            margin_pct = round((p_val / s_val) * 100, 2) if s_val > 0 else 0.0

            all_rows.append({
                "Order_ID": f"ORD-2024-{order_counter:05d}",
                "Order_Date": dt.strftime("%Y-%m-%d"),
                "Month": m["Month"],
                "Month_Num": m["MonthNum"],
                "Quarter": m["Quarter"],
                "Region": m_regions[i],
                "Category": cat,
                "Sub_Category": subcat,
                "Customer_Segment": m_segments[i],
                "Payment_Mode": m_payments[i],
                "Sales_INR": s_val,
                "Profit_INR": p_val,
                "Margin_Pct": margin_pct
            })
            order_counter += 1

    df_orders = pd.DataFrame(all_rows)

    # Adjust rounding discrepancy to match exact 7,20,00,000 and 1,65,00,000
    sales_diff = 72000000.0 - df_orders["Sales_INR"].sum()
    profit_diff = 16500000.0 - df_orders["Profit_INR"].sum()
    df_orders.loc[0, "Sales_INR"] = round(df_orders.loc[0, "Sales_INR"] + sales_diff, 2)
    df_orders.loc[0, "Profit_INR"] = round(df_orders.loc[0, "Profit_INR"] + profit_diff, 2)
    df_orders.loc[0, "Margin_Pct"] = round((df_orders.loc[0, "Profit_INR"] / df_orders.loc[0, "Sales_INR"]) * 100, 2)

    print(f"Total Orders: {len(df_orders):,}")
    print(f"Total Sales: Rs {df_orders['Sales_INR'].sum():,.2f} ({df_orders['Sales_INR'].sum()/1e7:.2f} Cr)")
    print(f"Total Profit: Rs {df_orders['Profit_INR'].sum():,.2f} ({df_orders['Profit_INR'].sum()/1e7:.2f} Cr)")
    print(f"Overall Profit Margin: {(df_orders['Profit_INR'].sum() / df_orders['Sales_INR'].sum()) * 100:.2f}%")

    # Export to CSV for Power BI
    csv_file = "Sales_Data_PowerBI.csv"
    df_orders.to_csv(csv_file, index=False)
    print(f"Saved: {csv_file}")

    # Create polished Excel workbook
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Create sheets
    ws_dash = wb.create_sheet(title="Dashboard")
    ws_dash.views.sheetView[0].showGridLines = False

    ws_summary = wb.create_sheet(title="Summary_Tables")
    ws_data = wb.create_sheet(title="Raw_Data")

    # Fonts & Colors
    font_family = "Segoe UI"
    navy_dark = "1B365D"
    card_bg = "1A365D"
    teal_color = "26A69A"
    coral_color = "E06D53"
    gray_bg = "F4F6F9"
    white_color = "FFFFFF"
    border_color = "D1D5DB"

    thin_border = Border(
        left=Side(style='thin', color=border_color),
        right=Side(style='thin', color=border_color),
        top=Side(style='thin', color=border_color),
        bottom=Side(style='thin', color=border_color)
    )

    header_font = Font(name=font_family, size=18, bold=True, color="1E293B")
    card_value_font = Font(name=font_family, size=18, bold=True, color="FFFFFF")
    card_label_font = Font(name=font_family, size=9, bold=False, color="CBD5E1")
    table_header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    table_body_font = Font(name=font_family, size=10, bold=False, color="1E293B")

    # 1. Populate Summary_Tables
    # Monthly Summary
    ws_summary["A1"] = "Month"
    ws_summary["B1"] = "Sales (Rs Lakh)"
    ws_summary["C1"] = "Profit (Rs Lakh)"
    ws_summary["D1"] = "Profit Margin %"
    ws_summary["E1"] = "Quarter"

    monthly_rows = [
        ("Jan", 42, 8, "=C2/B2", "Q1"),
        ("Feb", 45, 9, "=C3/B3", "Q1"),
        ("Mar", 51, 11, "=C4/B4", "Q1"),
        ("Apr", 48, 9, "=C5/B5", "Q2"),
        ("May", 55, 12, "=C6/B6", "Q2"),
        ("Jun", 60, 14, "=C7/B7", "Q2"),
        ("Jul", 58, 13, "=C8/B8", "Q3"),
        ("Aug", 63, 15, "=C9/B9", "Q3"),
        ("Sep", 67, 16, "=C10/B10", "Q3"),
        ("Oct", 70, 17, "=C11/B11", "Q4"),
        ("Nov", 76, 19, "=C12/B12", "Q4"),
        ("Dec", 85, 22, "=C13/B13", "Q4"),
    ]
    for r_idx, r_data in enumerate(monthly_rows, start=2):
        for c_idx, val in enumerate(r_data, start=1):
            cell = ws_summary.cell(row=r_idx, column=c_idx, value=val)
            if c_idx in [2, 3]:
                cell.number_format = "#,##0"
            elif c_idx == 4:
                cell.number_format = "0.0%"

    # Quarterly Summary
    ws_summary["G1"] = "Quarter"
    ws_summary["H1"] = "Quarterly Sales (Rs Lakh)"
    ws_summary["I1"] = "Quarterly Profit (Rs Lakh)"
    q_rows = [
        ("Q1", 138, 28),
        ("Q2", 163, 35),
        ("Q3", 188, 44),
        ("Q4", 231, 58),
    ]
    for r_idx, r_data in enumerate(q_rows, start=2):
        for c_idx, val in enumerate(r_data, start=7):
            cell = ws_summary.cell(row=r_idx, column=c_idx, value=val)
            if c_idx in [8, 9]:
                cell.number_format = "#,##0"

    # Regional Summary
    ws_summary["K1"] = "Region"
    ws_summary["L1"] = "Sales Share %"
    ws_summary["M1"] = "Sales (Rs Cr)"
    reg_rows = [
        ("West", 0.38, 2.736),
        ("South", 0.27, 1.944),
        ("North", 0.20, 1.440),
        ("East", 0.15, 1.080),
    ]
    for r_idx, r_data in enumerate(reg_rows, start=2):
        for c_idx, val in enumerate(r_data, start=11):
            cell = ws_summary.cell(row=r_idx, column=c_idx, value=val)
            if c_idx == 12:
                cell.number_format = "0%"
            elif c_idx == 13:
                cell.number_format = "0.00"

    # Category Summary (Order matching the chart: Toys at top or Electronics)
    ws_summary["O1"] = "Category"
    ws_summary["P1"] = "Sales (Rs Cr)"
    cat_rows = [
        ("Electronics", 2.4),
        ("Furniture", 1.7),
        ("Clothing", 1.3),
        ("Grocery", 1.1),
        ("Toys", 0.7),
    ]
    for r_idx, r_data in enumerate(cat_rows, start=2):
        for c_idx, val in enumerate(r_data, start=15):
            cell = ws_summary.cell(row=r_idx, column=c_idx, value=val)
            if c_idx == 16:
                cell.number_format = "0.0"

    # Style summary table headers
    for col_letter in ["A", "B", "C", "D", "E", "G", "H", "I", "K", "L", "M", "O", "P"]:
        c = ws_summary[f"{col_letter}1"]
        c.font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")

    # 2. Populate Dashboard Sheet
    ws_dash.row_dimensions[1].height = 20
    ws_dash.row_dimensions[2].height = 36
    ws_dash.row_dimensions[3].height = 15
    ws_dash.row_dimensions[4].height = 28
    ws_dash.row_dimensions[5].height = 20
    ws_dash.row_dimensions[6].height = 20

    # Title
    ws_dash["B2"] = "Sample Sales Performance Dashboard"
    ws_dash["B2"].font = header_font
    ws_dash["B2"].alignment = Alignment(horizontal="center", vertical="center")
    ws_dash.merge_cells("B2:P2")

    # KPI Cards Setup
    # Card 1: B4:D5 -> Total Sales (Rs 7.2 Cr)
    # Card 2: F4:H5 -> Total Profit (Rs 1.65 Cr)
    # Card 3: J4:L5 -> Profit Margin (22.9%)
    # Card 4: N4:P5 -> Total Orders (18,450)

    cards = [
        {"val_cell": "B4", "val_range": "B4:D4", "lbl_cell": "B5", "lbl_range": "B5:D5", "val": "Rs 7.2 Cr", "label": "Total Sales"},
        {"val_cell": "F4", "val_range": "F4:H4", "lbl_cell": "F5", "lbl_range": "F5:H5", "val": "Rs 1.65 Cr", "label": "Total Profit"},
        {"val_cell": "J4", "val_range": "J4:L4", "lbl_cell": "J5", "lbl_range": "J5:L5", "val": "22.9%", "label": "Profit Margin"},
        {"val_cell": "N4", "val_range": "N4:P4", "lbl_cell": "N5", "lbl_range": "N5:P5", "val": "18,450", "label": "Total Orders"},
    ]

    fill_card = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")

    for card in cards:
        ws_dash.merge_cells(card["val_range"])
        ws_dash.merge_cells(card["lbl_range"])
        v_cell = ws_dash[card["val_cell"]]
        v_cell.value = card["val"]
        v_cell.font = card_value_font
        v_cell.alignment = Alignment(horizontal="center", vertical="center")

        l_cell = ws_dash[card["lbl_cell"]]
        l_cell.value = card["label"]
        l_cell.font = card_label_font
        l_cell.alignment = Alignment(horizontal="center", vertical="center")

        # Color the whole card block
        col_start = card["val_cell"][0]
        col_end = card["val_range"].split(":")[1][0]
        c_start_num = openpyxl.utils.column_index_from_string(col_start)
        c_end_num = openpyxl.utils.column_index_from_string(col_end)
        for r in [4, 5]:
            for c in range(c_start_num, c_end_num + 1):
                ws_dash.cell(row=r, column=c).fill = fill_card

    # Add Charts
    # Chart 1: Line Chart - Monthly Sales vs Profit
    chart1 = LineChart()
    chart1.title = "Monthly Sales vs Profit (Rs Lakh)"
    chart1.style = 13
    chart1.y_axis.title = ""
    chart1.x_axis.title = ""
    chart1.width = 16
    chart1.height = 8.5

    data1 = Reference(ws_summary, min_col=2, min_row=1, max_col=3, max_row=13)
    cats1 = Reference(ws_summary, min_col=1, min_row=2, max_row=13)
    chart1.add_data(data1, titles_from_data=True)
    chart1.set_categories(cats1)

    # Style series lines
    s1 = chart1.series[0]
    s1.graphicalProperties.line.solidFill = "1B365D"
    s1.graphicalProperties.line.width = 25000
    s1.marker.symbol = "circle"
    s1.marker.graphicalProperties.solidFill = "1B365D"
    s1.marker.graphicalProperties.line.solidFill = "1B365D"

    s2 = chart1.series[1]
    s2.graphicalProperties.line.solidFill = "26A69A"
    s2.graphicalProperties.line.width = 25000
    s2.marker.symbol = "circle"
    s2.marker.graphicalProperties.solidFill = "26A69A"
    s2.marker.graphicalProperties.line.solidFill = "26A69A"

    ws_dash.add_chart(chart1, "B7")

    # Chart 2: Pie Chart - Sales by Region
    chart2 = PieChart()
    chart2.title = "Sales by Region"
    chart2.width = 10.5
    chart2.height = 8.5
    data2 = Reference(ws_summary, min_col=12, min_row=1, max_row=5)
    cats2 = Reference(ws_summary, min_col=11, min_row=2, max_row=5)
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats2)
    chart2.dataLabels = DataLabelList()
    chart2.dataLabels.showPercent = True
    chart2.dataLabels.showVal = False
    chart2.dataLabels.showCatName = True
    ws_dash.add_chart(chart2, "K7")

    # Chart 3: Horizontal Bar Chart - Sales by Category
    chart3 = BarChart()
    chart3.type = "bar"
    chart3.title = "Sales by Category (Rs Cr)"
    chart3.style = 10
    chart3.y_axis.title = ""
    chart3.x_axis.title = ""
    chart3.width = 16
    chart3.height = 8.5
    chart3.legend = None

    data3 = Reference(ws_summary, min_col=16, min_row=1, max_row=6)
    cats3 = Reference(ws_summary, min_col=15, min_row=2, max_row=6)
    chart3.add_data(data3, titles_from_data=True)
    chart3.set_categories(cats3)
    if len(chart3.series) > 0:
        chart3.series[0].graphicalProperties.solidFill = "26A69A"
    ws_dash.add_chart(chart3, "B25")

    # Chart 4: Vertical Column Chart - Quarterly Sales
    chart4 = BarChart()
    chart4.type = "col"
    chart4.title = "Quarterly Sales (Rs Lakh)"
    chart4.style = 10
    chart4.y_axis.title = ""
    chart4.x_axis.title = ""
    chart4.width = 10.5
    chart4.height = 8.5
    chart4.legend = None

    data4 = Reference(ws_summary, min_col=8, min_row=1, max_row=5)
    cats4 = Reference(ws_summary, min_col=7, min_row=2, max_row=5)
    chart4.add_data(data4, titles_from_data=True)
    chart4.set_categories(cats4)
    if len(chart4.series) > 0:
        chart4.series[0].graphicalProperties.solidFill = "E06D53"
    ws_dash.add_chart(chart4, "K25")

    # 3. Populate Raw_Data Sheet (all 18,450 rows)
    headers = list(df_orders.columns)
    ws_data.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws_data.cell(row=1, column=col_num)
        cell.font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    print("Writing transactions to Excel Raw_Data sheet...")
    # Convert dataframe to tuples for fast bulk insertion
    for row in df_orders.itertuples(index=False):
        ws_data.append(list(row))

    # Format Raw_Data columns
    for row_idx in range(2, min(len(df_orders) + 2, 200)): # Format first 200 rows for preview speed
        ws_data.cell(row=row_idx, column=11).number_format = "#,##0.00"
        ws_data.cell(row=row_idx, column=12).number_format = "#,##0.00"
        ws_data.cell(row=row_idx, column=13).number_format = "0.00%"

    # Adjust column widths
    for sheet in [ws_summary, ws_data]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col[:50])
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    excel_file = "Sales_Performance_Dashboard.xlsx"
    wb.save(excel_file)
    print(f"Saved: {excel_file}")

if __name__ == "__main__":
    generate_dashboard_files()
